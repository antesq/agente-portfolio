"""Geração da planilha de portfólio em .xlsx (openpyxl), no formato MELI.

Abas geradas:
  - "Resumo": contagem de projetos por Status Geral.
  - "Contratação Meli": a tabela principal (mesmas colunas da planilha original).
  - "Plano de Ações": ações consolidadas geradas pelo agente.
"""
from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import COLUNAS_PLANILHA, STATUS_VALIDOS, Portfolio

# --- Paleta e estilos -----------------------------------------------------
_AZUL = "1F4E78"
_AZUL_CLARO = "D9E1F2"

_FILL_TITULO = PatternFill("solid", fgColor=_AZUL)
_FILL_CABECALHO = PatternFill("solid", fgColor=_AZUL_CLARO)
_FONT_TITULO = Font(bold=True, color="FFFFFF", size=13)
_FONT_CABECALHO = Font(bold=True, color=_AZUL, size=10)
_BORDA_FINA = Side(style="thin", color="BFBFBF")
_BORDA = Border(left=_BORDA_FINA, right=_BORDA_FINA, top=_BORDA_FINA, bottom=_BORDA_FINA)
_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Cores por Status Geral / prioridade.
_CORES = {
    "dentro do prazo": "C6EFCE",
    "concluído": "C6EFCE",
    "concluida": "C6EFCE",
    "baixa": "C6EFCE",
    "em risco": "FFEB9C",
    "congelado": "DDEBF7",
    "média": "FFEB9C",
    "em atraso": "FFC7CE",
    "cancelado": "F2F2F2",
    "alta": "FFC7CE",
    "urgente": "FFC7CE",
}


def gerar_planilha(portfolio: Portfolio, output_dir: Path) -> Path:
    """Gera o arquivo .xlsx e retorna o caminho salvo."""
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    _aba_resumo(wb.active, portfolio)
    _aba_contratacao(wb, portfolio)
    _aba_acoes(wb, portfolio)

    carimbo = datetime.now().strftime("%Y-%m-%d_%H%M")
    caminho = output_dir / f"portfolio_meli_{carimbo}.xlsx"
    wb.save(caminho)
    return caminho


# --------------------------------------------------------------------------
def _aba_resumo(ws: Worksheet, portfolio: Portfolio) -> None:
    ws.title = "Resumo"
    _titulo(ws, portfolio.titulo, colunas=3)
    ws.append([])
    ws.append([f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    ws.append([f"Total de projetos: {len(portfolio.projetos)}"])
    ws.append([])

    _linha_cabecalho(ws, ["Status Geral", "Qtd. Projetos", "% do Portfólio"])
    contagem = Counter(p.status_geral.strip() or "(sem status)" for p in portfolio.projetos)
    total = max(len(portfolio.projetos), 1)
    # Mostra primeiro os status conhecidos na ordem canônica, depois o resto.
    ordem = STATUS_VALIDOS + [s for s in contagem if s not in STATUS_VALIDOS]
    for status in ordem:
        if status in contagem:
            qtd = contagem[status]
            ws.append([status, qtd, f"{qtd / total * 100:.0f}%"])
            _colorir(ws.cell(row=ws.max_row, column=1), status)

    ws.append([])
    _linha_cabecalho(ws, ["Ações pendentes", "Qtd."])
    pendentes = sum(1 for a in portfolio.acoes if a.status.lower() != "concluída")
    ws.append(["Total de ações no plano", len(portfolio.acoes)])
    ws.append(["Ações ainda pendentes", pendentes])

    for i, largura in enumerate([28, 16, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura


def _aba_contratacao(wb: Workbook, portfolio: Portfolio) -> None:
    ws = wb.create_sheet(title="Contratação Meli")
    titulos = [titulo for _, titulo in COLUNAS_PLANILHA]
    _titulo(ws, "Contratação Meli", colunas=len(titulos))
    ws.append([])
    _linha_cabecalho(ws, titulos)

    idx_status = [campo for campo, _ in COLUNAS_PLANILHA].index("status_geral") + 1
    for p in portfolio.projetos:
        ws.append([getattr(p, campo) for campo, _ in COLUNAS_PLANILHA])
        _colorir(ws.cell(row=ws.max_row, column=idx_status), p.status_geral)

    # Larguras: escopo e observações largos, o resto médio.
    larguras = []
    for campo, _ in COLUNAS_PLANILHA:
        if campo in ("escopo_atual", "observacoes", "motivo_bloqueio"):
            larguras.append(45)
        elif campo in ("proposta_tecnica_contratada",):
            larguras.append(24)
        else:
            larguras.append(16)
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = _WRAP
    ws.freeze_panes = "B4"  # trava cabeçalho e a coluna de código


def _aba_acoes(wb: Workbook, portfolio: Portfolio) -> None:
    ws = wb.create_sheet(title="Plano de Ações")
    cabecalhos = ["Projeto", "Ação", "Origem", "Responsável", "Prazo", "Prioridade", "Status"]
    _titulo(ws, "Plano de Ações", colunas=len(cabecalhos))
    ws.append([])
    _linha_cabecalho(ws, cabecalhos)
    for a in portfolio.acoes:
        ws.append([a.projeto, a.acao, a.origem, a.responsavel, a.prazo, a.prioridade, a.status])
        _colorir(ws.cell(row=ws.max_row, column=6), a.prioridade)
        _colorir(ws.cell(row=ws.max_row, column=7), a.status)
    for i, largura in enumerate([14, 55, 16, 22, 16, 14, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = _WRAP
    ws.freeze_panes = "A4"


# --------------------------------------------------------------------------
def _titulo(ws: Worksheet, texto: str, colunas: int) -> None:
    ws.append([texto])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(colunas, 1))
    c = ws.cell(row=1, column=1)
    c.fill = _FILL_TITULO
    c.font = _FONT_TITULO
    c.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 26


def _linha_cabecalho(ws: Worksheet, cabecalhos: list[str]) -> None:
    ws.append(cabecalhos)
    linha = ws.max_row
    for col in range(1, len(cabecalhos) + 1):
        c = ws.cell(row=linha, column=col)
        c.fill = _FILL_CABECALHO
        c.font = _FONT_CABECALHO
        c.border = _BORDA
        c.alignment = _CENTRO


def _colorir(celula, valor: str) -> None:
    cor = _CORES.get(str(valor).strip().lower())
    if cor:
        celula.fill = PatternFill("solid", fgColor=cor)
