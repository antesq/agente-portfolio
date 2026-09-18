"""Visão de Carregamentos dos projetos BR (workspace 'BR Projetos' do Smartsheet).

Puxa, via API, a data da tarefa 'Carregamento' do cronograma de cada projeto do
workspace 'BR Projetos', classifica por urgência e publica um resumo (Adaptive
Card, estilo BI) no canal do Teams. Também exporta um Excel com o detalhe.
"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from integrations import teams_notifier

WS_BR_PROJETOS = 4886249304549252
_API = "https://api.smartsheet.com/2.0"

# Categorias de urgência (ordem de exibição) + emoji.
_CATEGORIAS = ["Atrasado", "Esta semana", "Próx. 2 semanas", "Futuro", "Sem data", "Concluído"]
_EMOJI = {
    "Atrasado": "🔴", "Esta semana": "🟡", "Próx. 2 semanas": "🟢",
    "Futuro": "⚪", "Sem data": "⬜", "Concluído": "✅",
}


def _headers() -> dict[str, str]:
    return {"Authorization": f"Bearer {config.SMARTSHEET_ACCESS_TOKEN}"}


def listar_sheets_br() -> list[tuple[int, str]]:
    """Lista (id, nome) de todos os sheets do workspace 'BR Projetos'."""
    r = requests.get(f"{_API}/workspaces/{WS_BR_PROJETOS}?loadAll=true",
                     headers=_headers(), timeout=60)
    r.raise_for_status()
    ws = r.json()
    out: list[tuple[int, str]] = []

    def _rec(node: dict) -> None:
        for s in node.get("sheets", []):
            out.append((s["id"], s["name"]))
        for f in node.get("folders", []):
            _rec(f)

    _rec(ws)
    return out


def _parse_nome(nome: str) -> dict[str, str]:
    """Extrai código do projeto e unidade a partir do nome do sheet."""
    # Remove um prefixo entre colchetes (ex.: "[ATUALIZAR] ") antes de ler o código.
    sem_prefixo = re.sub(r"^\s*\[[^\]]*\]\s*", "", nome)
    codigo = ""
    m = re.match(r"\s*(\d{3,4})", sem_prefixo)  # código = número no INÍCIO do nome
    if m:
        codigo = m.group(1)
    unidade = ""
    mu = re.search(r"\b((?:BR|X)?[A-Z]{2,4}\d{1,3})\b", nome)
    if mu:
        unidade = mu.group(1)
    return {"codigo": codigo, "unidade": unidade}


def _to_date(v: Any) -> date | None:
    if not v:
        return None
    try:
        return datetime.fromisoformat(str(v).replace("Z", "")).date()
    except ValueError:
        return None


def _pct_concluido(pct: Any) -> bool:
    return str(pct or "").strip().startswith("100")


def extrair_carregamento(sheet_id: int) -> dict[str, Any] | None:
    """Lê o cronograma e retorna os dados da tarefa 'Carregamento', se existir."""
    r = requests.get(f"{_API}/sheets/{sheet_id}", headers=_headers(), timeout=60)
    if r.status_code != 200:
        return None
    sh = r.json()
    colmap = {c["id"]: c["title"] for c in sh.get("columns", [])}
    for row in sh.get("rows", []):
        cells = {colmap.get(c["columnId"]): c.get("displayValue", c.get("value"))
                 for c in row.get("cells", [])}
        nome = str(cells.get("Nome da tarefa") or "").strip().lower()
        if nome == "carregamento":  # exato, para não pegar "Descarregamento"
            ini = _to_date(cells.get("Iniciar"))
            fim = _to_date(cells.get("Terminar"))
            return {
                "data": ini or fim,
                "inicio": ini,
                "fim": fim,
                "pct": cells.get("% concluído"),
                "concluido": _pct_concluido(cells.get("% concluído")),
                "responsavel": cells.get("Atribuído a") or "",
            }
    return None


def _classificar(data: date | None, concluido: bool, hoje: date) -> str:
    if concluido:
        return "Concluído"
    if data is None:
        return "Sem data"
    dias = (data - hoje).days
    if dias < 0:
        return "Atrasado"
    if dias <= 7:
        return "Esta semana"
    if dias <= 14:
        return "Próx. 2 semanas"
    return "Futuro"


def montar_visao() -> list[dict[str, Any]]:
    """Lê todos os projetos BR e devolve a lista de carregamentos classificada."""
    hoje = date.today()
    linhas: list[dict[str, Any]] = []
    for sid, nome in listar_sheets_br():
        info = _parse_nome(nome)
        if not info["codigo"]:
            continue  # ignora sheets que não são projetos (sem código)
        carga = extrair_carregamento(sid)
        data = carga["data"] if carga else None
        concluido = carga["concluido"] if carga else False
        linhas.append({
            "codigo": info["codigo"],
            "unidade": info["unidade"],
            "projeto": nome.strip(),
            "data_carregamento": data,
            "data_fmt": data.strftime("%d/%m/%Y") if data else "",
            "pct": (carga or {}).get("pct") or "",
            "responsavel": (carga or {}).get("responsavel") or "",
            "categoria": _classificar(data, concluido, hoje),
            "dias": (data - hoje).days if data else None,
        })
    # Ordena por data (sem data por último).
    linhas.sort(key=lambda x: (x["data_carregamento"] is None, x["data_carregamento"] or date.max))
    return linhas


# --------------------------------------------------------------------------
# Exportação Excel
# --------------------------------------------------------------------------
_CORES = {"Atrasado": "FFC7CE", "Esta semana": "FFEB9C", "Próx. 2 semanas": "C6EFCE",
          "Futuro": "D9E1F2", "Sem data": "F2F2F2", "Concluído": "E2EFDA"}


def exportar_excel(linhas: list[dict[str, Any]], output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Carregamentos BR"
    ws.append(["Carregamentos Install — Projetos BR",
               f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    ws["A1"].font = Font(bold=True, size=13, color="1F4E78")
    ws.append([])
    cab = ["Código", "Unidade", "Projeto", "Data Carregamento", "Situação", "% concl.", "Responsável"]
    ws.append(cab)
    for c in range(1, len(cab) + 1):
        cell = ws.cell(row=3, column=c)
        cell.font = Font(bold=True, color="1F4E78")
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for ln in linhas:
        ws.append([ln["codigo"], ln["unidade"], ln["projeto"], ln["data_fmt"],
                   ln["categoria"], ln["pct"], ln["responsavel"]])
        cell = ws.cell(row=ws.max_row, column=5)
        cor = _CORES.get(ln["categoria"])
        if cor:
            cell.fill = PatternFill("solid", fgColor=cor)
    larguras = [10, 12, 45, 18, 16, 10, 28]
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A4"
    carimbo = datetime.now().strftime("%Y-%m-%d_%H%M")
    caminho = output_dir / f"carregamentos_br_{carimbo}.xlsx"
    wb.save(caminho)
    return caminho


# --------------------------------------------------------------------------
# Cartão do Teams (BI resumido)
# --------------------------------------------------------------------------
def montar_card(linhas: list[dict[str, Any]], max_itens: int = 18) -> dict[str, Any]:
    from collections import Counter
    cont = Counter(ln["categoria"] for ln in linhas)
    total = len(linhas)

    resumo = "  ·  ".join(
        f"{_EMOJI[c]} {c}: **{cont[c]}**" for c in _CATEGORIAS if cont.get(c)
    )

    # Foco: atrasados + próximos (esta semana + 2 semanas), por data.
    foco = [ln for ln in linhas if ln["categoria"] in ("Atrasado", "Esta semana", "Próx. 2 semanas")]
    corpo: list[dict[str, Any]] = [
        {"type": "TextBlock", "text": "🚚 Carregamentos Install — Projetos BR",
         "weight": "Bolder", "size": "Large", "wrap": True},
        {"type": "TextBlock", "text": f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}  ·  {total} projetos",
         "isSubtle": True, "spacing": "None", "wrap": True},
        {"type": "TextBlock", "text": resumo or "Sem dados.", "wrap": True, "spacing": "Small"},
        {"type": "TextBlock", "text": "**Atrasados e próximos:**", "wrap": True, "spacing": "Medium"},
    ]
    linhas_card = []
    for ln in foco[:max_itens]:
        rotulo = f"{ln['codigo']} {ln['unidade']}".strip()
        data = ln["data_fmt"] or "sem data"
        linhas_card.append({
            "type": "ColumnSet",
            "spacing": "Small",
            "columns": [
                {"type": "Column", "width": "auto", "items": [
                    {"type": "TextBlock", "text": _EMOJI[ln["categoria"]]}]},
                {"type": "Column", "width": 40, "items": [
                    {"type": "TextBlock", "text": data, "weight": "Bolder"}]},
                {"type": "Column", "width": "stretch", "items": [
                    {"type": "TextBlock", "text": rotulo or ln["projeto"][:30], "wrap": True}]},
            ],
        })
    if not foco:
        linhas_card.append({"type": "TextBlock", "text": "Nenhum carregamento atrasado ou nas próximas 2 semanas. 🎉", "wrap": True})
    corpo.extend(linhas_card)
    if len(foco) > max_itens:
        corpo.append({"type": "TextBlock",
                      "text": f"...e mais {len(foco) - max_itens}. Veja o Excel completo.",
                      "isSubtle": True, "wrap": True})

    return {
        "type": "AdaptiveCard",
        "$schema": "http://adaptivecards.io/schemas/adaptive-card.json",
        "version": "1.4",
        "body": corpo,
    }


def gerar_e_publicar(publicar: bool = True) -> dict[str, Any]:
    """Fluxo completo: lê BR Projetos, exporta Excel e publica o card no Teams."""
    linhas = montar_visao()
    caminho_excel = exportar_excel(linhas, config.OUTPUT_DIR)
    resultado: dict[str, Any] = {
        "total_projetos": len(linhas),
        "excel": str(caminho_excel),
        "linhas": linhas,
    }
    if publicar:
        card = montar_card(linhas)
        resultado["teams"] = teams_notifier.publicar_card(card)
    return resultado
