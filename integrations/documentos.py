"""Leitura de documentos de proposta (PDF e XLSX) da pasta do projeto.

Contexto: quando chega um e-mail de proposta, os anexos ficam na pasta do
projeto no drive Z:, por exemplo:
    Z:\\PROJETOS\\2026\\Mercado Livre\\Projeto 1677 Barueri\\
        PPT-BR-2026-07-1704_01 - TÉCNICA.pdf      (proposta técnica)
        RELATÓRIO CUSTOS ENGENHARIA.pdf
        RELATÓRIO CUSTOS OBRA.pdf / .xlsx
        RELATÓRIO CUSTOS PRODUÇÃO E MATERIAL.pdf / .xlsx

Este módulo:
  - lista os documentos de uma pasta de projeto;
  - lê PDFs via a visão/PDF nativa do Claude e extrai campos estruturados;
  - lê os XLSX de custo com openpyxl para pegar números exatos.
"""
from __future__ import annotations

import base64
import json
import re
from pathlib import Path
from typing import Any

import anthropic
from openpyxl import load_workbook

import config

# Padrão do número de proposta técnica, ex.: PPT-BR-2026-07-1704_01
RE_PROPOSTA = re.compile(r"PPT-[A-Z]{2}-\d{4}-\d{2}-\d{3,4}(?:_\d{1,2})?", re.IGNORECASE)


def listar_documentos(pasta: str | Path) -> list[dict[str, Any]]:
    """Lista PDFs e planilhas de uma pasta de projeto."""
    p = Path(pasta)
    if not p.exists():
        raise RuntimeError(f"Pasta não encontrada: {pasta}")
    docs = []
    for arq in sorted(p.iterdir()):
        if arq.suffix.lower() in (".pdf", ".xlsx", ".xls"):
            docs.append({
                "nome": arq.name,
                "caminho": str(arq),
                "tipo": arq.suffix.lower().lstrip("."),
                "tamanho_kb": round(arq.stat().st_size / 1024),
            })
    return docs


def _ler_xlsx_resumo(caminho: Path, max_linhas: int = 60) -> str:
    """Extrai texto tabular de um XLSX de custo (para o Claude interpretar)."""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    partes: list[str] = []
    for ws in wb.worksheets:
        partes.append(f"### Aba: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_linhas:
                partes.append("... (linhas adicionais omitidas)")
                break
            valores = [str(c) for c in row if c is not None]
            if valores:
                partes.append(" | ".join(valores))
    wb.close()
    return "\n".join(partes)


def extrair_dados_proposta(pasta: str | Path) -> dict[str, Any]:
    """Lê os documentos da pasta e extrai os dados da proposta com o Claude.

    Retorna um dicionário com campos úteis para preencher o portfólio:
    proposta_tecnica, escopo_atual, cidade, unidade, custos, etc.
    """
    if not config.ANTHROPIC_API_KEY:
        raise RuntimeError("ANTHROPIC_API_KEY não configurado — necessário para ler documentos.")

    pasta = Path(pasta)
    docs = listar_documentos(pasta)
    if not docs:
        raise RuntimeError(f"Nenhum PDF/XLSX encontrado em: {pasta}")

    conteudo: list[dict[str, Any]] = []

    # PDFs entram como documentos nativos (o Claude lê texto e layout).
    for d in docs:
        if d["tipo"] == "pdf":
            dados = base64.standard_b64encode(Path(d["caminho"]).read_bytes()).decode("utf-8")
            conteudo.append({"type": "text", "text": f"Documento: {d['nome']}"})
            conteudo.append({
                "type": "document",
                "source": {"type": "base64", "media_type": "application/pdf", "data": dados},
            })

    # XLSX de custo entram como texto tabular (números exatos).
    for d in docs:
        if d["tipo"] in ("xlsx", "xls"):
            try:
                resumo = _ler_xlsx_resumo(Path(d["caminho"]))
                conteudo.append({"type": "text", "text": f"Planilha {d['nome']}:\n{resumo}"})
            except Exception as e:  # planilha corrompida/protegida não deve travar tudo
                conteudo.append({"type": "text", "text": f"(Falha ao ler {d['nome']}: {e})"})

    instrucao = (
        "Você recebeu os documentos de uma PROPOSTA de projeto de esteiras/transporte "
        "(cliente Mercado Livre / Install). Extraia e devolva SOMENTE um JSON com estas chaves:\n"
        '{\n'
        '  "proposta_tecnica_contratada": "código PPT-BR-... da proposta técnica",\n'
        '  "unidade_meli": "código da unidade, ex. SSP58",\n'
        '  "cidade": "cidade da obra",\n'
        '  "escopo_atual": "descrição resumida do escopo técnico (linhas, esteiras, metragens, fingers, etc.)",\n'
        '  "custo_engenharia": "valor total de engenharia, se houver",\n'
        '  "custo_obra": "valor total de obra, se houver",\n'
        '  "custo_producao_material": "valor total de produção e material, se houver",\n'
        '  "custo_total": "soma dos custos, se der para calcular",\n'
        '  "observacoes": "qualquer alerta relevante (revisão, pendência, divergência)"\n'
        '}\n'
        "Se um campo não estiver nos documentos, use string vazia. Não invente valores. "
        "Responda apenas o JSON, sem texto antes ou depois."
    )
    conteudo.append({"type": "text", "text": instrucao})

    client = anthropic.Anthropic(api_key=config.ANTHROPIC_API_KEY)
    resposta = client.messages.create(
        model=config.CLAUDE_MODEL,
        max_tokens=4000,
        messages=[{"role": "user", "content": conteudo}],
    )
    texto = "".join(b.text for b in resposta.content if b.type == "text").strip()

    dados = _json_seguro(texto)
    # Rede de segurança: se o Claude não achou o código da proposta, tenta pelo nome do arquivo.
    if not dados.get("proposta_tecnica_contratada"):
        for d in docs:
            m = RE_PROPOSTA.search(d["nome"])
            if m:
                dados["proposta_tecnica_contratada"] = m.group(0)
                break
    dados["_documentos_lidos"] = [d["nome"] for d in docs]
    return dados


def _json_seguro(texto: str) -> dict[str, Any]:
    """Extrai o primeiro bloco JSON de uma string, tolerando cercas de código."""
    texto = texto.strip()
    if texto.startswith("```"):
        texto = texto.strip("`")
        texto = texto[texto.find("{"):]
    inicio, fim = texto.find("{"), texto.rfind("}")
    if inicio == -1 or fim == -1:
        return {}
    try:
        return json.loads(texto[inicio: fim + 1])
    except json.JSONDecodeError:
        return {}
